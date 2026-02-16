"""Louisiana sports betting scraper — LSP Gaming Enforcement Division Excel reports.

Downloads the latest mobile and retail sports wagering Excel reports from
lsp.org. Each file contains FY sheets (FY22–FY26+) with monthly aggregate data.
Mobile launched January 2022; retail launched October 2021.

GGR = Net Proceeds (Wagers Written − Payouts to bettors).
Sports breakdown is GGR by sport (Net Proceeds by Sport/Type).
"""

import io
import re
from collections import defaultdict

import openpyxl
import pandas as pd
import requests
from bs4 import BeautifulSoup

from .base import BaseScraper

BASE_URL = "https://lsp.org"
CURRENT_URL = "https://lsp.org/about/leadershipsections/bureau-of-investigations/gaming-enforcement-division/gaming-revenue-reports/"
ARCHIVE_URL = "https://lsp.org/about/leadershipsections/bureau-of-investigations/gaming-enforcement-division/gaming-revenue-reports/gaming-revenue-reports-archive/"

# FY months: July (month 0) through June (month 11)
FY_MONTH_MAP = [7, 8, 9, 10, 11, 12, 1, 2, 3, 4, 5, 6]

# Sport columns in the Excel: column index -> canonical sport name
SPORT_COLUMNS = {
    9: "Baseball",
    10: "Basketball",
    11: "Football",
    12: "Soccer",
    13: "Parlay",
    14: "Other",
}


class LAScraper(BaseScraper):
    STATE_CODE = "la"
    STATE_NAME = "Louisiana"
    TAX_RATE = 0.15
    TAX_RATE_NOTE = "15% mobile, 10% retail on net proceeds"
    FREQUENCY = "monthly"
    HAS_OPERATOR_DATA = False
    SOURCE_URL = CURRENT_URL
    DATE_COLUMN = "Month"
    LAUNCH_DATE = "2021-10-31"
    GGR_DEFINITION = "Net Proceeds (Handle − Payouts to bettors)"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._mobile_wb = None
        self._retail_wb = None

    def scrape(self) -> dict[str, pd.DataFrame]:
        mobile_url, retail_url = self._find_latest_excel_urls()

        mobile_data = {}
        retail_data = {}

        if mobile_url:
            print(f"  Mobile: {mobile_url}")
            self._mobile_wb = self._download_workbook(mobile_url)
            if self._mobile_wb:
                mobile_data = self._parse_workbook(self._mobile_wb)
        else:
            print("  No mobile Excel URL found.")

        if retail_url:
            print(f"  Retail: {retail_url}")
            self._retail_wb = self._download_workbook(retail_url)
            if self._retail_wb:
                retail_data = self._parse_workbook(self._retail_wb)
        else:
            print("  No retail Excel URL found.")

        if not mobile_data and not retail_data:
            print("  No data parsed.")
            return {}

        # Combine mobile + retail by month
        all_dates = sorted(set(list(mobile_data.keys()) + list(retail_data.keys())))
        records = []
        for date_str in all_dates:
            m = mobile_data.get(date_str, {})
            r = retail_data.get(date_str, {})
            handle = (m.get("handle", 0) or 0) + (r.get("handle", 0) or 0)
            ggr = (m.get("ggr", 0) or 0) + (r.get("ggr", 0) or 0)
            if handle > 0:
                records.append({
                    "date": pd.Timestamp(date_str),
                    "handle": handle,
                    "ggr": ggr,
                })

        if not records:
            print("  No combined data.")
            return {}

        df = pd.DataFrame(records).sort_values("date").reset_index(drop=True)
        print(f"  Parsed {len(df)} months (combined mobile + retail).")
        return {"Total": df}

    def scrape_sports(self) -> dict[str, pd.DataFrame] | None:
        """Extract sports GGR breakdown from cached workbooks.

        Returns GGR by sport (Net Proceeds by Sport/Type) combined from
        mobile + retail. Uses 'handle' column name for framework compat.
        """
        if not self._mobile_wb and not self._retail_wb:
            return None

        mobile_sports = self._parse_sports(self._mobile_wb) if self._mobile_wb else {}
        retail_sports = self._parse_sports(self._retail_wb) if self._retail_wb else {}

        # Combine mobile + retail by sport by date
        all_sports = sorted(set(list(mobile_sports.keys()) + list(retail_sports.keys())))
        if not all_sports:
            return None

        result = {}
        for sport in all_sports:
            m_dates = {r["date"]: r["value"] for r in mobile_sports.get(sport, [])}
            r_dates = {r["date"]: r["value"] for r in retail_sports.get(sport, [])}
            all_dates = sorted(set(list(m_dates.keys()) + list(r_dates.keys())))

            records = []
            for d in all_dates:
                val = (m_dates.get(d, 0) or 0) + (r_dates.get(d, 0) or 0)
                records.append({"date": pd.Timestamp(d), "handle": val})

            if records:
                result[sport] = pd.DataFrame(records).sort_values("date").reset_index(drop=True)

        return result if result else None

    def _find_latest_excel_urls(self) -> tuple[str | None, str | None]:
        """Find the latest mobile and retail Excel URLs from the current page."""
        mobile_url = None
        retail_url = None

        try:
            resp = requests.get(CURRENT_URL, timeout=30)
            resp.raise_for_status()
            soup = BeautifulSoup(resp.text, "lxml")

            for a in soup.find_all("a", href=True):
                href = a["href"]
                if not href.lower().endswith(".xlsx"):
                    continue

                full_url = href if href.startswith("http") else BASE_URL + href
                lower = href.lower()

                # Take the first (most recent) mobile and retail Excel links
                if "mobile" in lower and "sb" in lower and not mobile_url:
                    mobile_url = full_url
                elif "retail" in lower and "sb" in lower and not retail_url:
                    retail_url = full_url

                if mobile_url and retail_url:
                    break

        except Exception as e:
            print(f"  Error fetching page: {e}")

        return mobile_url, retail_url

    def _download_workbook(self, url: str):
        """Download an Excel file and return an openpyxl workbook."""
        try:
            resp = requests.get(url, timeout=60)
            if resp.status_code != 200:
                print(f"    HTTP {resp.status_code}")
                return None
            return openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
        except Exception as e:
            print(f"    Download error: {e}")
            return None

    def _parse_workbook(self, wb) -> dict[str, dict]:
        """Parse handle/GGR from all FY sheets. Returns {date_str: {handle, ggr}}."""
        data = {}
        for sheet_name in wb.sheetnames:
            fy_match = re.match(r"FY(\d{2})", sheet_name)
            if not fy_match:
                continue

            fy_num = int(fy_match.group(1))
            fy_start_year = 2000 + fy_num - 1  # FY22 starts July 2021

            ws = wb[sheet_name]
            for month_offset in range(12):
                row_num = month_offset + 2  # Row 1 = header, data starts row 2

                handle = ws.cell(row=row_num, column=3).value  # Column C: Wagers Written
                ggr = ws.cell(row=row_num, column=5).value  # Column E: Net Proceeds

                if not handle or handle <= 0:
                    continue

                # Compute correct date from FY position (avoid cell date errors)
                cal_month = FY_MONTH_MAP[month_offset]
                cal_year = fy_start_year if cal_month >= 7 else fy_start_year + 1
                date_str = f"{cal_year}-{cal_month:02d}-01"

                data[date_str] = {"handle": handle, "ggr": ggr if ggr else 0}

        print(f"    Parsed {len(data)} months from {len(wb.sheetnames)} sheets")
        return data

    def _parse_sports(self, wb) -> dict[str, list]:
        """Parse sports GGR breakdown from all FY sheets.

        Returns {sport_name: [{date, value}, ...]}
        """
        sports = defaultdict(list)

        for sheet_name in wb.sheetnames:
            fy_match = re.match(r"FY(\d{2})", sheet_name)
            if not fy_match:
                continue

            fy_num = int(fy_match.group(1))
            fy_start_year = 2000 + fy_num - 1

            ws = wb[sheet_name]
            for month_offset in range(12):
                row_num = month_offset + 2
                handle = ws.cell(row=row_num, column=3).value
                if not handle or handle <= 0:
                    continue

                cal_month = FY_MONTH_MAP[month_offset]
                cal_year = fy_start_year if cal_month >= 7 else fy_start_year + 1
                date_str = f"{cal_year}-{cal_month:02d}-01"

                for col, sport_name in SPORT_COLUMNS.items():
                    val = ws.cell(row=row_num, column=col).value
                    if val is not None:
                        sports[sport_name].append({"date": date_str, "value": val})

        return dict(sports)
